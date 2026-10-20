import io
import re

import openpyxl

from finance_tracker.storage.excel import _build_s3_key, create_excel_bytes


def _workbook(expenses, incomes):
    return openpyxl.load_workbook(io.BytesIO(create_excel_bytes(expenses, incomes)))


def _expense(amount, currency="COP", category="comida"):
    return {
        "amount": amount,
        "currency": currency,
        "category": category,
        "expense_date": "2026-10-20",
        "payment_method": None,
        "merchant": None,
        "description": "x",
        "source": "text",
    }


def _income(amount, currency="COP", category="salario"):
    return {
        "amount": amount,
        "currency": currency,
        "category": category,
        "income_date": "2026-10-15",
        "description": "y",
        "source": "manual",
    }


def test_three_sheets_with_spanish_headers():
    wb = _workbook([_expense(10000)], [_income(50000)])

    assert wb.sheetnames == ["Gastos", "Ingresos", "Resumen"]
    assert [c.value for c in wb["Gastos"][1]][:4] == [
        "Monto",
        "Moneda",
        "Categoría",
        "Fecha",
    ]
    assert [c.value for c in wb["Ingresos"][1]][:4] == [
        "Monto",
        "Moneda",
        "Categoría",
        "Fecha",
    ]
    assert (
        wb["Gastos"]["A2"].value == 10000 and wb["Gastos"]["D2"].value == "2026-10-20"
    )
    assert (
        wb["Ingresos"]["A2"].value == 50000
        and wb["Ingresos"]["D2"].value == "2026-10-15"
    )


def test_summary_has_totals_and_net_per_currency():
    wb = _workbook(
        [_expense(10000), _expense(5000), _expense(20, "USD")],
        [_income(50000), _income(100, "USD")],
    )
    rows = [[c.value for c in row] for row in wb["Resumen"].iter_rows()]
    assert rows == [
        ["Moneda", "Total gastos", "Total ingresos", "Neto"],
        ["COP", 15000, 50000, 35000],
        ["USD", 20, 100, 80],
    ]


def test_empty_export_still_has_every_sheet():
    wb = _workbook([], [])
    assert wb.sheetnames == ["Gastos", "Ingresos", "Resumen"]
    assert wb["Resumen"].max_row == 1


def test_s3_key_uses_bogota_time():
    key = _build_s3_key("+573001234567", "2026-10-01", "2026-10-22")
    assert re.fullmatch(
        r"processed/whatsapp/573001234567/\d{8}T\d{6}_2026-10-01_2026-10-22\.xlsx", key
    )
